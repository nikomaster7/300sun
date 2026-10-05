# Rebuilds docs/agencias/300sun-comisiones-agencias.xlsx (empty, with two example rows). Careful: it overwrites the file.
# python3 tools/make-commission-sheet.py   (needs: pip install openpyxl)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from datetime import date
F='Arial'
H=Font(name=F,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1A1A1A')
BLUE=Font(name=F,color='0000FF'); BLACK=Font(name=F); YEL=PatternFill('solid',fgColor='FFF2CC')
EUR='#,##0.00 €;-#,##0.00 €;-'
wb=Workbook()
aj=wb.active; aj.title='Ajustes'
aj['A1']='Ajustes'; aj['A1'].font=Font(name=F,bold=True,size=14)
aj['A2']='Comisión para agencias'; aj['B2']=0.15; aj['B2'].number_format='0%'; aj['B2'].font=BLUE; aj['B2'].fill=YEL
aj['B2'].comment=Comment('Decidido por Nicolás el 2026-10-06 (15 %). Si lo cambias aquí, cambia también 300sun.com/advisors y el PDF.','300sun')
aj['A3']='Días para pagar la comisión'; aj['B3']=7; aj['B3'].font=BLUE; aj['B3'].fill=YEL
aj['B3'].comment=Comment('Lo que promete 300sun.com/advisors: pago en los 7 días siguientes al paseo.','300sun')
aj['A4']='Enlace base para agencias'; aj['B4']='https://300sun.com/cruise/?ref='; aj['B4'].font=BLUE
for c in ('A2','A3','A4'): aj[c].font=BLACK
aj['A6']='Cómo usar esta hoja'; aj['A6'].font=Font(name=F,bold=True)
notes=['Rellena solo las celdas en azul (con fondo amarillo las más importantes). Lo negro se calcula solo.',
 'Agencias: una fila por agencia. El "código ref" va en minúsculas y con guiones (ej. sunny-cruises); el enlace sale solo.',
 'Reservas: una fila por paseo o paella vendido por una agencia. Pon lo que pagó el cliente (sin propinas, entradas ni comida).',
 'Cuando pagues la comisión, cambia "Pagada" a Sí y pon fecha y método. La hoja Agencias te dice cuánto debes a cada una.',
 'Las filas de "test-agency" son un ejemplo: bórralas cuando tengas reservas reales.']
for i,n in enumerate(notes): aj.cell(row=7+i,column=1,value='• '+n).font=BLACK
aj.column_dimensions['A'].width=34; aj.column_dimensions['B'].width=36

rs=wb.create_sheet('Reservas')
hdr=['Fecha del paseo','Código ref (agencia)','Agencia','Cliente','Servicio','Personas','Importe cobrado al cliente (€)','Comisión %','Comisión (€)','Pagar antes de','Pagada','Fecha de pago','Método','Notas']
for i,h in enumerate(hdr,1):
    c=rs.cell(row=1,column=i,value=h); c.font=H; c.fill=HF; c.alignment=Alignment(vertical='center',wrap_text=True)
rs.row_dimensions[1].height=32
N=301
for r in range(2,N+1):
    rs.cell(row=r,column=1).number_format='DD/MM/YYYY'
    rs.cell(row=r,column=3,value=f'=IF(B{r}="","",IFERROR(INDEX(Agencias!$A$2:$A$32,MATCH(B{r},Agencias!$B$2:$B$32,0)),"¿Agencia nueva?"))')
    rs.cell(row=r,column=7).number_format=EUR
    rs.cell(row=r,column=8,value=f'=IF(G{r}="","",Ajustes!$B$2)'); rs.cell(row=r,column=8).number_format='0%'
    rs.cell(row=r,column=9,value=f'=IF(G{r}="","",ROUND(G{r}*H{r},2))'); rs.cell(row=r,column=9).number_format=EUR
    rs.cell(row=r,column=10,value=f'=IF(A{r}="","",A{r}+Ajustes!$B$3)'); rs.cell(row=r,column=10).number_format='DD/MM/YYYY'
    rs.cell(row=r,column=12).number_format='DD/MM/YYYY'
    for col in (1,2,4,5,6,7,11,12,13,14): rs.cell(row=r,column=col).font=BLUE
    for col in (3,8,9,10): rs.cell(row=r,column=col).font=BLACK
rows=[[date(2026,11,20),'test-agency',None,'Smith family (ejemplo)','Paseo + personas extra',5,180,None,None,None,'Sí',date(2026,11,24),'Wise','Ejemplo: 150 € + 1 persona extra 30 €'],
      [date(2026,12,3),'test-agency',None,'Jones couple (ejemplo)','Paseo privado 3 h',2,150,None,None,None,'No',None,None,'Ejemplo: comisión aún sin pagar']]
for k,row in enumerate(rows):
    for i,v in enumerate(row,1):
        if v is not None: rs.cell(row=2+k,column=i,value=v)
for rng,lst in ((f'K2:K{N}','"Sí,No"'),(f'M2:M{N}','"PayPal,Wise,Transferencia"'),(f'E2:E{N}','"Paseo privado 3 h,Paseo + personas extra,Taller de paella,Paseo + paella,Grupo / alojamiento"')):
    dv=DataValidation(type='list',formula1=lst,allow_blank=True); rs.add_data_validation(dv); dv.add(rng)
for L,w in zip('ABCDEFGHIJKLMN',(13,20,22,24,22,10,18,11,13,14,9,13,13,40)): rs.column_dimensions[L].width=w
rs.freeze_panes='A2'
rs['G1'].comment=Comment('Lo que pagó el cliente por el paseo o la paella (señal + resto). Sin propinas, entradas ni comida.','300sun')

ag=wb.create_sheet('Agencias')
hdr=['Agencia','Código ref','Enlace para la agencia','Contacto','Email / WhatsApp','Total vendido (€)','Comisión total (€)','Comisión pendiente (€)']
for i,h in enumerate(hdr,1):
    c=ag.cell(row=1,column=i,value=h); c.font=H; c.fill=HF; c.alignment=Alignment(vertical='center',wrap_text=True)
ag.row_dimensions[1].height=30
for r in range(2,33):
    ag.cell(row=r,column=3,value=f'=IF(B{r}="","",Ajustes!$B$4&B{r})').font=BLACK
    ag.cell(row=r,column=6,value=f'=IF(B{r}="","",SUMIFS(Reservas!$G$2:$G${N},Reservas!$B$2:$B${N},B{r}))')
    ag.cell(row=r,column=7,value=f'=IF(B{r}="","",SUMIFS(Reservas!$I$2:$I${N},Reservas!$B$2:$B${N},B{r}))')
    ag.cell(row=r,column=8,value=f'=IF(B{r}="","",SUMIFS(Reservas!$I$2:$I${N},Reservas!$B$2:$B${N},B{r},Reservas!$K$2:$K${N},"No"))')
    for col in (6,7,8): ag.cell(row=r,column=col).number_format=EUR; ag.cell(row=r,column=col).font=BLACK
    for col in (1,2,4,5): ag.cell(row=r,column=col).font=BLUE
for i,v in enumerate(['Agencia de prueba','test-agency',None,'(nombre del advisor)','advisor@example.com'],1):
    if v is not None: ag.cell(row=2,column=i,value=v)
ag['A34']='TOTAL'; ag['A34'].font=Font(name=F,bold=True)
for col,L in ((6,'F'),(7,'G'),(8,'H')):
    c=ag.cell(row=34,column=col,value=f'=SUM({L}2:{L}33)'); c.number_format=EUR; c.font=Font(name=F,bold=True)
for L,w in zip('ABCDEFGH',(24,18,44,22,26,16,16,18)): ag.column_dimensions[L].width=w
ag.freeze_panes='A2'
wb.move_sheet('Ajustes', offset=2)
wb.active=0
from openpyxl.workbook.properties import CalcProperties
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save('docs/agencias/300sun-comisiones-agencias.xlsx')
